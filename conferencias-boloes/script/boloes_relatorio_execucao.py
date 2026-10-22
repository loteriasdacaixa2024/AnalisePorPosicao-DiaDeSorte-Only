# -*- coding: utf-8 -*-
"""
Relatórios da execução automática: resumo por modalidade, relatório geral,
Excel profissional e JSON consolidado (sem apagar os arquivos individuais).
"""
from __future__ import annotations

import os
from collections import Counter
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

from boloes_consolidar import carregar_json_boloes, salvar_json_boloes

STATUS_OK = ('CONCLUÍDO', 'CONCLUIDO', 'VAZIO', 'EXTRAÇÃO COMPLETA', 'EXTRACAO COMPLETA')
STATUS_ERRO = ('ERRO',)
STATUS_PEND = ('PENDENTE', 'PAUSADO', 'EM ANDAMENTO', 'AGUARDANDO')


def fmt_duracao(segundos: Any) -> str:
    try:
        secs = max(0, int(segundos))
    except (TypeError, ValueError):
        return '00h 00min 00s'
    h, rem = divmod(secs, 3600)
    m, s = divmod(rem, 60)
    return f'{h:02d}h {m:02d}min {s:02d}s'


def fmt_tamanho(bytes_n: Any, *, mb: bool = False) -> str:
    try:
        n = float(bytes_n or 0)
    except (TypeError, ValueError):
        return '0.0 KB'
    if mb or n >= 1024 * 1024:
        return f'{n / (1024 * 1024):.2f} MB'
    return f'{n / 1024:.1f} KB'


def _parse_dt(txt: str) -> Optional[datetime]:
    raw = (txt or '').strip()
    if not raw:
        return None
    for fmt in ('%Y-%m-%d %H:%M:%S', '%d/%m/%Y %H:%M:%S', '%Y-%m-%dT%H:%M:%S'):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def _fmt_br(dt: Optional[datetime], fallback: str = '') -> str:
    if dt is None:
        return fallback or '—'
    return dt.strftime('%d/%m/%Y %H:%M:%S')


def segundos_entre(inicio: str, fim: str) -> int:
    t0 = _parse_dt(inicio)
    t1 = _parse_dt(fim)
    if not t0 or not t1:
        return 0
    return max(0, int((t1 - t0).total_seconds()))


def stats_arquivo_json(pasta_json: str, arquivo_base: str) -> Tuple[int, int, List[str], str]:
    """Retorna (registros, bytes, ufs, caminho)."""
    base = (arquivo_base or '').removesuffix('.json').strip()
    path = os.path.join(pasta_json, f'{base}.json') if base else ''
    if not path or not os.path.isfile(path):
        return 0, 0, [], path
    boloes = carregar_json_boloes(path)
    ufs = sorted({
        str(b.get('uf') or '').upper()
        for b in boloes
        if str(b.get('uf') or '').strip()
    })
    try:
        tamanho = os.path.getsize(path)
    except OSError:
        tamanho = 0
    return len(boloes), tamanho, ufs, path


def coletar_linhas_modalidades(
    ciclo: Dict[str, Any],
    pasta_json: str,
    ordem_slugs: Optional[List[str]] = None,
) -> List[Dict[str, Any]]:
    mods = ciclo.get('modalidades') or {}
    slugs = list(ordem_slugs) if ordem_slugs else list(mods.keys())
    linhas: List[Dict[str, Any]] = []
    for slug in slugs:
        mod = mods.get(slug) or {}
        ab = str(mod.get('arquivo_base') or '').removesuffix('.json')
        n, tam, ufs, path = stats_arquivo_json(pasta_json, ab)
        if not n:
            n = int(mod.get('registros') or (mod.get('ufs') or {}).get('LISTA', {}).get('boloes') or 0)
        if not ufs and mod.get('ufs_encontradas'):
            ufs = [u.strip() for u in str(mod.get('ufs_encontradas')).split(',') if u.strip()]
        inicio = mod.get('inicio') or ''
        fim = mod.get('fim') or ''
        st = str(mod.get('status') or '—')
        linhas.append({
            'slug': slug,
            'modalidade': mod.get('label') or slug,
            'concurso': str(mod.get('concurso') or ''),
            'arquivo': f'{ab}.json' if ab else '—',
            'arquivo_base': ab,
            'path': path,
            'registros': n,
            'bytes': tam,
            'tamanho': fmt_tamanho(tam, mb=tam >= 1024 * 1024),
            'estados': ', '.join(ufs),
            'ufs': ufs,
            'inicio': inicio,
            'fim': fim,
            'tempo_seg': segundos_entre(inicio, fim),
            'tempo': fmt_duracao(segundos_entre(inicio, fim)),
            'status': st,
            'erro': str((mod.get('ufs') or {}).get('LISTA', {}).get('erro') or ''),
        })
    return linhas


def status_geral_de(linhas: List[Dict[str, Any]], pausado: bool = False) -> str:
    if pausado:
        return 'PAUSADO'
    if not linhas:
        return 'CONCLUÍDO'
    statuses = [str(l.get('status') or '').upper() for l in linhas]
    if any(s in STATUS_ERRO or s in STATUS_PEND for s in statuses):
        if any(s in STATUS_OK for s in statuses) or any(s in STATUS_ERRO for s in statuses):
            return 'CONCLUÍDO COM OCORRÊNCIAS'
        return 'CONCLUÍDO COM OCORRÊNCIAS'
    return 'CONCLUÍDO'


def imprimir_resumo_modalidade(linha: Dict[str, Any], out_fn=print) -> None:
    out_fn('')
    out_fn('=' * 60)
    out_fn('MODALIDADE CONCLUÍDA')
    out_fn('=' * 60)
    out_fn('')
    out_fn(f"Modalidade: {linha.get('modalidade') or '—'}")
    out_fn(f"Concurso: {linha.get('concurso') or '—'}")
    out_fn('')
    out_fn('Arquivo:')
    out_fn(str(linha.get('arquivo') or '—'))
    out_fn('')
    out_fn('Registros extraídos:')
    out_fn(str(linha.get('registros') or 0))
    out_fn('')
    out_fn('Tamanho:')
    out_fn(str(linha.get('tamanho') or '—'))
    out_fn('')
    out_fn('Estados extraídos:')
    out_fn(str(linha.get('estados') or '—'))
    out_fn('')
    out_fn('Tempo da modalidade:')
    out_fn(str(linha.get('tempo') or '—'))
    out_fn('')
    out_fn('Status:')
    out_fn(str(linha.get('status') or '—'))
    if linha.get('erro'):
        out_fn('')
        out_fn(f"Ocorrência: {linha['erro']}")
    out_fn('=' * 60)


def imprimir_relatorio_geral(
    ciclo: Dict[str, Any],
    linhas: List[Dict[str, Any]],
    *,
    status_geral: str = '',
    out_fn=print,
) -> None:
    inicio = ciclo.get('inicio') or ''
    fim = ciclo.get('fim') or ''
    t0 = _parse_dt(inicio)
    t1 = _parse_dt(fim) or datetime.now()
    tempo = fmt_duracao(segundos_entre(inicio, fim) if fim else segundos_entre(inicio, t1.strftime('%Y-%m-%d %H:%M:%S')))
    if not tempo or (not fim and t0):
        tempo = fmt_duracao(int((t1 - t0).total_seconds()) if t0 else 0)

    total_reg = sum(int(l.get('registros') or 0) for l in linhas)
    total_bytes = sum(int(l.get('bytes') or 0) for l in linhas)
    st = status_geral or status_geral_de(linhas, str(ciclo.get('status') or '') == 'PAUSADO')

    out_fn('')
    out_fn('=' * 60)
    out_fn('RELATÓRIO FINAL DA EXTRAÇÃO')
    out_fn('=' * 60)
    out_fn('')
    out_fn('Início:')
    out_fn(_fmt_br(t0, inicio or '—'))
    out_fn('')
    out_fn('Fim:')
    out_fn(_fmt_br(t1 if fim else datetime.now(), fim or '—'))
    out_fn('')
    out_fn('Tempo total:')
    out_fn(tempo)
    out_fn('')
    out_fn('')
    out_fn('MODALIDADES')
    out_fn('-' * 60)
    for i, l in enumerate(linhas, start=1):
        out_fn('')
        out_fn(f"{i}. {l.get('modalidade') or l.get('slug')}")
        out_fn(f"   Concurso: {l.get('concurso') or '—'}")
        out_fn(f"   Registros: {l.get('registros') or 0}")
        out_fn(f"   Arquivo: {l.get('arquivo') or '—'}")
        out_fn(f"   Tamanho: {l.get('tamanho') or '—'}")
        out_fn(f"   Tempo: {l.get('tempo') or '—'}")
        out_fn(f"   Estados: {l.get('estados') or '—'}")
        out_fn(f"   Status: {l.get('status') or '—'}")
        if l.get('erro'):
            out_fn(f"   Ocorrência: {l['erro']}")
    out_fn('')
    out_fn('=' * 60)
    out_fn('TOTAL GERAL')
    out_fn('=' * 60)
    out_fn('')
    out_fn(f'Modalidades processadas: {len(linhas)}')
    out_fn('')
    out_fn(f'Registros totais: {total_reg}')
    out_fn('')
    out_fn('Tamanho total dos JSON:')
    out_fn(fmt_tamanho(total_bytes, mb=True))
    out_fn('')
    out_fn('Tempo total:')
    out_fn(tempo)
    out_fn('')
    out_fn('Status geral:')
    out_fn(st)
    out_fn('=' * 60)


def _linhas_estados(linhas: List[Dict[str, Any]], pasta_json: str) -> List[Dict[str, Any]]:
    saida: List[Dict[str, Any]] = []
    for l in linhas:
        path = l.get('path') or ''
        boloes = carregar_json_boloes(path) if path else []
        cont: Counter = Counter()
        for b in boloes:
            uf = str(b.get('uf') or '').upper().strip()
            if uf:
                cont[uf] += 1
        if not cont and l.get('ufs'):
            for uf in l['ufs']:
                cont[uf] = 0
        for uf, qtd in sorted(cont.items()):
            saida.append({
                'modalidade': l.get('modalidade'),
                'concurso': l.get('concurso'),
                'uf': uf,
                'registros': qtd,
            })
    return saida


def _linhas_detalhe(ciclo: Dict[str, Any], linhas: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    det: List[Dict[str, Any]] = []
    det.append({
        'quando': ciclo.get('inicio') or '',
        'modalidade': 'CICLO',
        'evento': 'INÍCIO',
        'detalhe': f"id={ciclo.get('ciclo_id') or '—'}",
        'status': ciclo.get('status') or '',
    })
    for l in linhas:
        det.append({
            'quando': l.get('inicio') or '',
            'modalidade': l.get('modalidade'),
            'evento': 'MODALIDADE',
            'detalhe': (
                f"conc={l.get('concurso') or '—'} | "
                f"{l.get('registros') or 0} reg. | "
                f"{l.get('arquivo')} | UFs={l.get('estados') or '—'}"
            ),
            'status': l.get('status') or '',
        })
        if l.get('erro'):
            det.append({
                'quando': l.get('fim') or l.get('inicio') or '',
                'modalidade': l.get('modalidade'),
                'evento': 'ERRO',
                'detalhe': l.get('erro'),
                'status': 'ERRO',
            })
    det.append({
        'quando': ciclo.get('fim') or datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'modalidade': 'CICLO',
        'evento': 'FIM',
        'detalhe': f"status={ciclo.get('status') or '—'}",
        'status': ciclo.get('status') or '',
    })
    return det


def _estilo_status(ws, col: int, linhas_n: int, fill_map: dict, Font, PatternFill, Alignment) -> None:
    for row in range(4, 4 + linhas_n):
        cell = ws.cell(row=row, column=col)
        val = str(cell.value or '').upper()
        fill = None
        if any(s in val for s in ('CONCLUÍDO', 'CONCLUIDO', 'OK', 'COMPLETA')) and 'OCORR' not in val:
            fill = fill_map['ok']
        elif 'OCORR' in val or 'PEND' in val or 'PAUS' in val:
            fill = fill_map['warn']
        elif 'ERRO' in val:
            fill = fill_map['erro']
        if fill:
            cell.fill = fill
            cell.font = Font(name='Calibri', size=11, bold=True)
        cell.alignment = Alignment(horizontal='center')


def gerar_excel_relatorio(
    pasta_json: str,
    ciclo: Dict[str, Any],
    linhas: List[Dict[str, Any]],
    *,
    agora: Optional[datetime] = None,
) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    agora = agora or datetime.now()
    nome = f"relatorio_extracao_{agora.strftime('%Y-%m-%d_%H-%M')}.xlsx"
    path = os.path.join(pasta_json, nome)

    wb = Workbook()
    thin = Border(
        left=Side(style='thin', color='B0B7C3'),
        right=Side(style='thin', color='B0B7C3'),
        top=Side(style='thin', color='B0B7C3'),
        bottom=Side(style='thin', color='B0B7C3'),
    )
    fill_titulo = PatternFill('solid', fgColor='1F4E79')
    fill_header = PatternFill('solid', fgColor='2E75B6')
    fill_zebra = PatternFill('solid', fgColor='F2F5F8')
    fill_ok = PatternFill('solid', fgColor='C6EFCE')
    fill_warn = PatternFill('solid', fgColor='FFF2CC')
    fill_erro = PatternFill('solid', fgColor='F8CBAD')
    font_titulo = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
    font_sub = Font(name='Calibri', size=11, italic=True, color='D6E3F0')
    font_head = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    font_cell = Font(name='Calibri', size=11)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_center = Alignment(horizontal='center', vertical='center')

    def _montar_aba(ws, titulo: str, subtitulo: str, headers: List[str], rows: List[list], *, status_col: Optional[int] = None):
        n_cols = len(headers)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
        c1 = ws.cell(1, 1, titulo)
        c1.font = font_titulo
        c1.fill = fill_titulo
        c1.alignment = Alignment(horizontal='left', vertical='center')
        c2 = ws.cell(2, 1, subtitulo)
        c2.font = font_sub
        c2.fill = fill_titulo
        c2.alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 18
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(3, col, h)
            cell.font = font_head
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin
        for i, row in enumerate(rows):
            excel_row = 4 + i
            for col, val in enumerate(row, start=1):
                cell = ws.cell(excel_row, col, val)
                cell.font = font_cell
                cell.border = thin
                cell.alignment = align_left
                if i % 2 == 1:
                    cell.fill = fill_zebra
        if rows:
            ws.auto_filter.ref = f"A3:{get_column_letter(n_cols)}{3 + len(rows)}"
        ws.freeze_panes = 'A4'
        ws.auto_filter.ref = f"A3:{get_column_letter(n_cols)}{max(3, 3 + len(rows))}"
        for col in range(1, n_cols + 1):
            letter = get_column_letter(col)
            maxlen = len(str(headers[col - 1]))
            for row in rows:
                maxlen = max(maxlen, len(str(row[col - 1] if col - 1 < len(row) else '')))
            ws.column_dimensions[letter].width = min(48, max(12, maxlen + 3))
        if status_col and rows:
            _estilo_status(
                ws, status_col, len(rows),
                {'ok': fill_ok, 'warn': fill_warn, 'erro': fill_erro},
                Font, PatternFill, Alignment,
            )
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddHeader.left.text = titulo
        ws.print_title_rows = '1:3'

    st_geral = status_geral_de(linhas, str(ciclo.get('status') or '') == 'PAUSADO')
    sub = (
        f"Ciclo {ciclo.get('ciclo_id') or '—'}  |  "
        f"{_fmt_br(_parse_dt(ciclo.get('inicio') or ''))} → "
        f"{_fmt_br(_parse_dt(ciclo.get('fim') or '') or agora)}  |  "
        f"{st_geral}"
    )

    ws1 = wb.active
    ws1.title = 'Resumo Geral'
    rows1 = []
    for l in linhas:
        rows1.append([
            l.get('modalidade'),
            l.get('concurso') or '',
            l.get('arquivo'),
            int(l.get('registros') or 0),
            l.get('tamanho'),
            l.get('estados') or '',
            l.get('inicio') or '',
            l.get('fim') or '',
            l.get('tempo'),
            l.get('status'),
        ])
    _montar_aba(
        ws1, 'Resumo geral da extração', sub,
        ['Modalidade', 'Concurso', 'Arquivo', 'Registros', 'Tamanho',
         'Estados', 'Início', 'Fim', 'Tempo', 'Status'],
        rows1, status_col=10,
    )
    for row in range(4, 4 + len(rows1)):
        ws1.cell(row, 4).number_format = '#,##0'
        ws1.cell(row, 4).alignment = align_center

    ws2 = wb.create_sheet('Estados')
    est = _linhas_estados(linhas, pasta_json)
    rows2 = [[e['modalidade'], e['concurso'], e['uf'], int(e['registros'])] for e in est]
    _montar_aba(
        ws2, 'Registros por estado (UF)', sub,
        ['Modalidade', 'Concurso', 'UF', 'Quantidade de registros'],
        rows2,
    )
    for row in range(4, 4 + len(rows2)):
        ws2.cell(row, 4).number_format = '#,##0'
        ws2.cell(row, 3).alignment = align_center

    ws3 = wb.create_sheet('Detalhamento')
    det = _linhas_detalhe(ciclo, linhas)
    rows3 = [[d['quando'], d['modalidade'], d['evento'], d['detalhe'], d['status']] for d in det]
    _montar_aba(
        ws3, 'Detalhamento da execução', sub,
        ['Quando', 'Modalidade', 'Evento', 'Detalhe', 'Status'],
        rows3, status_col=5,
    )

    wb.save(path)
    return path


def gerar_json_consolidado(
    pasta_json: str,
    linhas: List[Dict[str, Any]],
    *,
    agora: Optional[datetime] = None,
) -> Tuple[str, int]:
    """
    Cria boloes_consolidado_YYYY-MM-DD.json sem apagar os JSON individuais.
    Metadados de origem são adicionados de forma não destrutiva.
    """
    agora = agora or datetime.now()
    nome = f"boloes_consolidado_{agora.strftime('%Y-%m-%d')}.json"
    path = os.path.join(pasta_json, nome)
    reunidos: List[dict] = []
    vistos = set()
    for l in linhas:
        src = l.get('path') or ''
        if not src or not os.path.isfile(src):
            continue
        origem_arquivo = os.path.basename(src)
        for b in carregar_json_boloes(src):
            if not isinstance(b, dict):
                continue
            h = b.get('hash_bolao') or ''
            chave = h or f"{origem_arquivo}:{b.get('id_bolao_api') or b.get('indice')}"
            if chave in vistos:
                continue
            vistos.add(chave)
            item = dict(b)
            if 'origem' not in item:
                item['origem'] = item.get('fonte') or 'api'
            item.setdefault('origem_arquivo', origem_arquivo)
            if not item.get('modalidade'):
                item['modalidade'] = l.get('modalidade')
            if not item.get('concurso'):
                item['concurso'] = l.get('concurso')
            if not item.get('estado') and item.get('uf'):
                item['estado'] = item.get('uf')
            reunidos.append(item)
    for i, b in enumerate(reunidos, start=1):
        b['indice_consolidado'] = i
    salvar_json_boloes(path, reunidos)
    return path, len(reunidos)


def gerar_entregaveis_finais(
    pasta_json: str,
    ciclo: Dict[str, Any],
    ordem_slugs: Optional[List[str]] = None,
    *,
    out_fn=print,
    pausado: bool = False,
) -> Dict[str, Any]:
    linhas = coletar_linhas_modalidades(ciclo, pasta_json, ordem_slugs)
    st = status_geral_de(linhas, pausado)
    ciclo['status_geral'] = st
    imprimir_relatorio_geral(ciclo, linhas, status_geral=st, out_fn=out_fn)

    excel = ''
    try:
        excel = gerar_excel_relatorio(pasta_json, ciclo, linhas)
        out_fn(f'\n  [EXCEL] Relatório: {os.path.basename(excel)}')
    except Exception as exc:
        out_fn(f'\n  [EXCEL] Falha ao gerar relatório: {exc}')

    consolidado = ''
    n_cons = 0
    try:
        consolidado, n_cons = gerar_json_consolidado(pasta_json, linhas)
        out_fn(f'  [JSON] Consolidado: {os.path.basename(consolidado)} ({n_cons} reg.)')
        out_fn('  [JSON] Arquivos individuais preservados.')
    except Exception as exc:
        out_fn(f'  [JSON] Falha ao consolidar: {exc}')

    return {
        'linhas': linhas,
        'status_geral': st,
        'excel': excel,
        'consolidado': consolidado,
        'registros_consolidados': n_cons,
    }
